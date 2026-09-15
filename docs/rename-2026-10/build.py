#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""master.json から 3点の資料を作り直す。

    python3 build.py

  出力先は out/ 。
    ・商品名リネーム（2026年10月 全店改変）.pdf   … 決裁用の4ページ
    ・商品名リネーム（2026年10月 全店改変）.xlsx  … 全55品の一覧
    ・HP更新用マスタ（2026年10月）.md              … 公式HPに渡す36品

  直すのは master.json だけ。このファイルは「見た目」を持っているだけなので、
  文言・商品・区分を変えたいときに開く必要はない。
"""
import json, os, sys

from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.utils import ImageReader

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'out')
IMG = os.path.join(HERE, 'images')

# ── 紙とフォント ────────────────────────────────────────────────
PW, PH = 595.2756, 841.8898          # A4 縦
M, R = 39.6850, 555.5906             # 左余白 14mm / 右端
CW = R - M                           # 本文の幅

FONT = 'JP'
FONT_CANDIDATES = [
    '/usr/share/fonts/opentype/ipafont-gothic/ipagp.ttf',   # IPA Pゴシック（元のPDFと同じ）
    '/usr/share/fonts/truetype/fonts-japanese-gothic.ttf',
    '/System/Library/Fonts/ヒラギノ角ゴシック W3.ttc',
    '/Library/Fonts/Arial Unicode.ttf',
]

# ── 色 ──────────────────────────────────────────────────────────
DARK   = (0.133333, 0.133333, 0.133333)   # #222222 文字・帯
MUTED  = (0.4, 0.4, 0.4)                  # #666666 補足
FAINT  = (0.866667, 0.866667, 0.866667)   # #DDDDDD 帯の中の補足
BEIGE  = (0.956863, 0.945098, 0.917647)   # #F4F1EA 面
WHITE  = (1, 1, 1)
RED    = (0.784314, 0.211765, 0.168627)   # #C8362B 変更するもの
GOLD   = (0.788235, 0.635294, 0.152941)   # #C9A227 印・罫
NAVY   = (0.184314, 0.282353, 0.345098)   # #2F4858 和名維持

BADGE = {                                  # 変更区分 → （ラベル, 色）
    '和名維持':     ('和名維持',   NAVY),
    '英語名変更':   ('英語名を変更', RED),
    '再販時に修正': ('再販時に修正', GOLD),
}
XLSX_FILL = {
    '和名維持': 'DCE6F1', '英語名変更': 'F8D7D3',
    '変更不要': 'F2F2F2', '再販時に修正': 'FFF2CC',
}

# ── ページ内の位置（元のPDFから実測。pt） ──────────────────────
BAR_H, GOLD_H = 62.3622, 2.8346            # 上の黒帯と金の線
TITLE_Y, SUB_Y, FOOT_Y = 34.0157, 51.0236, 819.2126
LEAD_Y, LEAD_LH = 94.5600, 13.7750         # リード文の1行目と行送り

CARD_X, CARD_W = M, CW
CARD_H_DESC, CARD_H_PLAIN = 85.0394, 65.1969
CARD_GAP_DESC, CARD_GAP_PLAIN = 91.3386 - 85.0394, 70.8661 - 65.1969
IMG_CX, IMG_BOTTOM_PAD = 96.3780, 8.5039   # 写真の中心と、カード下端からの余白
IMG_H_DESC, IMG_H_PLAIN = 68.0315, 48.1890
BADGE_X, BADGE_W, BADGE_H = 158.7402, 56.6929, 14.1732
NAME_X = 223.9370
JA_GAP = 10.4                              # 英語名と「／ 日本語名」の間
BOX_X, BOX_W, BOX_H = 158.7402, 385.5118, 28.3465
BOX_PAD_X = 170.0787 - 158.7402
BAR_W = 4.2520                             # 説明文の左の金の線

# カード上端からの相対位置
DY_GENRE, DY_BADGE, DY_NAME, DY_JA = 11.40, 11.94, 24.99, 24.46
DY_OLD, DY_NOTE, DY_FOOT = 44.34, 57.14, 77.00
DY_BOX, DY_BOX_NO_OLD = 49.14, 36.84
DY_BOX_LABEL, DY_BOX_BODY = 11.36, 23.60

TBL_PAD_X = 5.6693                         # 表の左右の内側余白
CALLOUT_X = M + 14.1732                    # 囲みの中の文字の左
WRAP_W = 490.0                             # 囲みの中の折り返し幅

ASCENT = 0.8795                            # 字上端 → ベースライン（IPAゴシック）


def _find_font():
    for p in FONT_CANDIDATES:
        if os.path.exists(p):
            return p
    sys.exit('日本語フォントが見つかりません。IPAゴシックを入れてください：\n'
             '  apt-get install fonts-ipafont-gothic')


pdfmetrics.registerFont(TTFont(FONT, _find_font()))


def w(s, size):
    return pdfmetrics.stringWidth(s, FONT, size)


class Sheet:
    """上端からの y（pt）で書ける薄いラッパー。"""

    def __init__(self, path, title):
        self.c = canvas.Canvas(path, pagesize=(PW, PH))
        self.c.setTitle(title)
        self.c.setAuthor('株式会社ANCHOR／Age.3')
        self.c.setSubject('商品名リネーム')

    # --- 描くもの -------------------------------------------------
    def text(self, x, y, s, size, color=DARK, bold=False, align='left'):
        if not s:
            return
        if align == 'right':
            x -= w(s, size)
        elif align == 'center':
            x -= w(s, size) / 2
        self.c.setFillColorRGB(*color)
        self.c.setFont(FONT, size)
        self.c.drawString(x, PH - y, s)
        if bold:                       # 太字を持たないフォントなので0.25ptずらして二度打ち
            self.c.drawString(x + 0.25, PH - y, s)

    def rect(self, x, y, width, height, color, radius=0):
        self.c.setFillColorRGB(*color)
        if radius:
            self.c.roundRect(x, PH - y - height, width, height, radius, stroke=0, fill=1)
        else:
            self.c.rect(x, PH - y - height, width, height, stroke=0, fill=1)

    def image(self, name, cx, bottom, height):
        path = os.path.join(IMG, name)
        if not os.path.exists(path):
            return
        img = ImageReader(path)
        iw, ih = img.getSize()
        width = iw * height / ih
        self.c.drawImage(img, cx - width / 2, PH - bottom, width, height, mask='auto')

    # --- ページ ---------------------------------------------------
    def page(self, meta, subtitle, num, total):
        self.rect(0, 0, PW, BAR_H, DARK)
        self.rect(0, BAR_H, PW, GOLD_H, GOLD)
        self.text(M, TITLE_Y, meta['title'], 14, WHITE, bold=True)
        self.text(M, SUB_Y, subtitle, 8.5, FAINT)
        self.text(R, SUB_Y, meta['header_right'], 8, FAINT, align='right')
        self.text(M, FOOT_Y, meta['footer'], 7.5, MUTED)
        self.text(R, FOOT_Y, '%d / %d' % (num, total), 7.5, MUTED, align='right')

    def done(self):
        self.c.showPage()

    def save(self):
        self.c.save()


def wrap(s, size, width):
    """空白で区切って折り返し、空白のない和文のかたまりは幅のところで割る。"""
    lines, cur = [], ''
    for word in s.split(' '):
        cand = word if not cur else cur + ' ' + word
        if w(cand, size) <= width:
            cur = cand
            continue
        if cur:
            lines.append(cur)
            cur = ''
        while w(word, size) > width:
            n = 1
            while n < len(word) and w(word[:n + 1], size) <= width:
                n += 1
            lines.append(word[:n])
            word = word[n:]
        cur = word
    if cur:
        lines.append(cur)
    return lines or ['']


# ── 1〜2ページ目・4ページ目のカード ────────────────────────────
def draw_card(s, top, it):
    has_desc = bool(it.get('desc_en'))
    h = CARD_H_DESC if has_desc else CARD_H_PLAIN
    s.rect(CARD_X, top, CARD_W, h, BEIGE, radius=8.5039)

    img_h = IMG_H_DESC if has_desc else IMG_H_PLAIN
    if it.get('image'):
        s.image(it['image'], IMG_CX, top + h - IMG_BOTTOM_PAD, img_h)

    s.text(R, top + DY_GENRE, 'ジャンル帯：' + it['genre'], 7, MUTED, align='right')

    label, color = BADGE[it['kind']]
    s.rect(BADGE_X, top + DY_BADGE, BADGE_W, BADGE_H, color, radius=2.8346)
    s.text(BADGE_X + BADGE_W / 2, top + DY_BADGE + 10.46, label, 7, WHITE,
           bold=True, align='center')

    ja = '／ ' + it['ja'] if it['ja'] else ''
    size = 12
    for cand in (12, 10, 9):             # 名前が長いときだけ落とす
        size = cand
        if NAME_X + w(it['en'], cand) + JA_GAP + w(ja, 9) <= R:
            break
    s.text(NAME_X, top + DY_NAME, it['en'], size, DARK, bold=True)
    s.text(NAME_X + w(it['en'], size) + JA_GAP, top + DY_JA, ja, 9, MUTED)

    old = it.get('old_en')
    if old and old != it['en']:
        s.text(BADGE_X, top + DY_OLD, '現行：' + old, 7.5, RED)

    if has_desc:
        box_top = top + (DY_BOX if old and old != it['en'] else DY_BOX_NO_OLD)
        s.rect(BOX_X, box_top, BOX_W, BOX_H, WHITE, radius=4.2520)
        s.rect(BOX_X, box_top, BAR_W, BOX_H, GOLD, radius=2)
        s.text(BOX_X + BOX_PAD_X, box_top + DY_BOX_LABEL, '説明文', 7, GOLD, bold=True)
        s.text(BOX_X + BOX_PAD_X, box_top + DY_BOX_BODY, it['desc_en'], 10, DARK, bold=True)
        if it.get('footnote'):
            s.text(BADGE_X, top + DY_FOOT, it['footnote'], 7.5, MUTED)
    elif it.get('note'):
        s.text(BADGE_X, top + DY_NOTE, it['note'], 8, MUTED)

    return h + (CARD_GAP_DESC if has_desc else CARD_GAP_PLAIN)


def draw_lead(s, y, lines, size=9.5, color=DARK):
    for i, line in enumerate(lines):
        s.text(M, y + i * LEAD_LH, line, size, color)
    # リード文の下に置くものの上端
    return y + (len(lines) - 1) * LEAD_LH + (12.73 if len(lines) > 1 else 10.33)


# ── 表 ──────────────────────────────────────────────────────────
def table(s, top, headers, cols, rows, head_h, head_dy, row_h, row_dy, head_size, cell):
    """cols=各列の左x／rows=各行の値／cell=列ごとの (文字の大きさ, 色, 太字)"""
    s.rect(M, top, CW, head_h, DARK)
    for x, label in zip(cols, headers):
        s.text(x, top + head_dy, label, head_size, WHITE, bold=True)
    y = top + head_h
    for i, row in enumerate(rows):
        if i % 2 == 0:
            s.rect(M, y, CW, row_h, BEIGE)
        for x, value, (size, color, bold) in zip(cols, row, cell):
            s.text(x, y + row_dy, value, size, color, bold=bold)
        y += row_h
    return y


def draw_box(s, top, title, body, pad=13.65, min_h=0):
    lines = wrap(body, 8, WRAP_W)
    h = max(min_h, pad + len(lines) * 11.6)
    s.rect(M, top, CW, h, BEIGE, radius=4.2520)
    s.rect(M, top, BAR_W, h, GOLD, radius=2)
    s.text(CALLOUT_X, top + 14.2, title, 9.5, DARK, bold=True)
    for i, line in enumerate(lines):
        s.text(CALLOUT_X, top + 27.8 + i * 11.6, line, 8, MUTED)
    return h


# ── PDF ─────────────────────────────────────────────────────────
def build_pdf(m, path):
    items = m['items']
    meta, sec = m['meta'], m['sections']
    wamei = [i for i in items if i['kind'] == '和名維持']
    rename = sorted([i for i in items if i['kind'] == '英語名変更'],
                    key=lambda i: i.get('card_seq', 999))
    wamei.sort(key=lambda i: i.get('card_seq', 999))
    keep = [i for i in items if i['kind'] == '変更不要']
    reprint = [i for i in items if i['kind'] == '再販時に修正']

    # 英語名変更のカードを2ページに割る（1ページ目は和名維持の下に入るだけ）
    first_page = 5
    pages_of_rename = [rename[:first_page], rename[first_page:]]
    pages_of_rename = [p for p in pages_of_rename if p]
    total = 2 + len(pages_of_rename)

    s = Sheet(path, meta['pdf_title'])

    # 1ページ目 ────────────────────────────────
    s.page(meta, meta['policy'], 1, total)
    y = draw_lead(s, LEAD_Y, sec['wamei_lead'])
    for it in wamei:
        y += draw_card(s, y, it)
    n = len(rename)
    lead = [l.format(n=n, i=1, t=len(pages_of_rename)) for l in sec['rename_lead']]
    y = draw_lead(s, y + 17.87, lead)
    for it in pages_of_rename[0]:
        y += draw_card(s, y, it)
    s.done()

    # 2ページ目以降（英語名変更の続き） ─────────
    for idx, group in enumerate(pages_of_rename[1:], start=2):
        s.page(meta, sec['page2_subtitle'], idx, total)
        lead = [l.format(n=n, i=idx, t=len(pages_of_rename)) for l in sec['rename_lead_cont']]
        y = draw_lead(s, LEAD_Y, lead)
        for it in group:
            y += draw_card(s, y, it)
        s.done()

    # 変更不要＋運用ルール ──────────────────────
    s.page(meta, sec['page3_subtitle'], total - 1, total)
    s.text(M, 90.7, sec['page3_table_title'], 10, DARK, bold=True)
    rows = [(i['en'], i['ja'], i['genre']) for i in keep]
    y = table(s, 104.8819, ['英語名', '日本語名', 'ジャンル帯'],
              [M + TBL_PAD_X, 235.7, 411.1], rows,
              head_h=14.1732, head_dy=10.18, row_h=12.4724, row_dy=9.88,
              head_size=7.5,
              cell=[(8, DARK, True), (7.5, MUTED, False), (7.5, MUTED, False)])
    y += 17.04
    for callout in m['callouts']:
        y += draw_box(s, y, callout['title'], callout['body'], min_h=28.3465) + 5.67
    y += 11.35 - 5.67
    s.text(M, y, m['rules_title'], 10, DARK, bold=True)
    y += 17.0
    for rule in m['rules']:
        y += draw_box(s, y, rule['title'], rule['body'], pad=25.5) + 5.67
    s.done()

    # 再販時に修正 ──────────────────────────────
    s.page(meta, sec['page4_subtitle'], total, total)
    s.rect(M, 85.0394, CW, 31.1811, RED, radius=4.2520)
    s.text(M + CW / 2, 105.74, sec['page4_banner'], 13, WHITE, bold=True, align='center')
    s.text(M, 136.5, sec['page4_lead'], 9, MUTED)

    cards = [i for i in reprint if i.get('image')]
    cards.sort(key=lambda i: i.get('card_seq', 999))
    y = 154.8
    for it in cards:
        y += draw_card(s, y, it)

    def chip(top, label, color):
        s.rect(M, top, 62.3622, 15.5906, color, radius=2.8346)
        s.text(M + 62.3622 / 2, top + 11.3, label, 7.5, WHITE, bold=True, align='center')

    wamei_rows = [i for i in reprint if i.get('desc_en')]
    other_rows = [i for i in reprint if i not in cards and i not in wamei_rows]
    other_rows += [i for i in items if i.get('reprint_row')]

    y += 15.6
    chip(y, '和名維持', NAVY)
    y += 24.1
    y = table(s, y, ['商品（日本語）', '英語名', '説明文'],
              [M + TBL_PAD_X, 141.7, 311.8],
              [(i['ja'], i['en'], i.get('desc_en', '')) for i in wamei_rows],
              head_h=15.5906, head_dy=11.34, row_h=17.0079, row_dy=11.8,
              head_size=8,
              cell=[(8, DARK, False), (8.5, DARK, True), (8, MUTED, False)])
    y += 16.4
    s.text(M, y, sec['page4_wamei_note'], 8, MUTED)

    y += 19.0
    chip(y, '英語名', RED)
    y += 24.1
    table(s, y, ['商品（日本語）', '英語名', '補足'],
          [M + TBL_PAD_X, 141.7, 311.8],
          [(i.get('reprint_row', {}).get('ja', i['ja']),
            i['en'],
            i.get('reprint_row', {}).get('note', i.get('note', ''))) for i in other_rows],
          head_h=15.5906, head_dy=11.34, row_h=17.0079, row_dy=11.8,
          head_size=8,
          cell=[(8, DARK, False), (8.5, DARK, True), (8, MUTED, False)])
    s.done()
    s.save()


# ── xlsx ────────────────────────────────────────────────────────
def build_xlsx(m, path):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '商品名マスタ'
    headers = ['No', '大分類', 'ジャンル帯', '英語名（新）', '日本語名',
               '現行英語名（旧）', '変更区分', '英語説明文', '備考']
    widths = [6, 9, 24, 36, 32, 30, 14, 42, 48]
    head_fill = PatternFill('solid', fgColor='222222')
    center = Alignment(vertical='center', wrap_text=True)

    ws.append(headers)
    for i, (cell, width) in enumerate(zip(ws[1], widths), start=1):
        cell.font = Font(bold=True, size=10, color='FFFFFF')
        cell.fill = head_fill
        cell.alignment = center
        ws.column_dimensions[chr(64 + i)].width = width
    ws.row_dimensions[1].height = 22

    for it in m['items']:
        ws.append([it['no'], it['group'], it['genre'], it['en'], it['ja'],
                   it.get('old_en', ''), it['kind'], it.get('desc_en', ''),
                   it.get('remark', '')])
        row = ws[ws.max_row]
        for cell in row:
            cell.font = Font(size=10)
            cell.alignment = center
        row[3].font = Font(size=10, bold=True)
        fill = XLSX_FILL.get(it['kind'])
        if fill:
            row[6].fill = PatternFill('solid', fgColor=fill)

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = 'A1:I%d' % ws.max_row

    ws2 = wb.create_sheet('凡例')
    ws2.column_dimensions['A'].width = 16
    ws2.column_dimensions['B'].width = 90
    for i, row in enumerate(m['legend'], start=1):
        ws2.append(row)
        for cell in ws2[i]:
            cell.alignment = center
            if i == 1:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = head_fill
        fill = XLSX_FILL.get(row[0])          # 変更区分の行は一覧と同じ色をつける
        if i > 1 and fill:
            ws2.cell(i, 1).fill = PatternFill('solid', fgColor=fill)
    wb.save(path)


# ── HP更新用の md ───────────────────────────────────────────────
def build_md(m, path):
    hp = m['hp']
    items = [i for i in m['items'] if i.get('hp')]
    items.sort(key=lambda i: i['no'])
    out = ['# ' + hp['title'], '']
    for line in hp['intro']:
        out.append(('  - ' if line.startswith('  ') else '- ') + line.strip())
    out += ['', '## 商品一覧（%d品）' % len(items), '',
            '| No | ジャンル帯 | 英語名（新） | 日本語名 | 旧英語名 | 変更区分 | '
            '英語説明文 | Web用ファイル名（zip） | Dropbox ファイル名（regular） |',
            '|---|---|---|---|---|---|---|---|---|']
    for i in items:
        out.append('| %d | %s | %s | %s | %s | %s | %s | %s | %s |' % (
            i['no'], i['genre'], i['en'], i['ja'], i.get('old_en', ''), i['kind'],
            i.get('desc_en', ''), i['hp']['web_filename'], i['hp']['dropbox_filename']))
    desc = [i for i in items if i.get('desc_en')]
    out += ['', '## 英語説明文を表示する商品', '', hp['desc_section'], '',
            '## 最終確認チェックリスト', '']
    out += ['- [ ] ' + c for c in hp['checklist']]
    out.append('')
    open(path, 'w', encoding='utf-8').write('\n'.join(out))
    return len(items), len(desc)


def main():
    m = json.load(open(os.path.join(HERE, 'master.json'), encoding='utf-8'))
    os.makedirs(OUT, exist_ok=True)
    meta = m['meta']

    pdf = os.path.join(OUT, meta['pdf_filename'])
    build_pdf(m, pdf)
    print('PDF   :', os.path.basename(pdf))

    xlsx = os.path.join(OUT, meta['xlsx_filename'])
    build_xlsx(m, xlsx)
    print('Excel :', os.path.basename(xlsx), '（%d品）' % len(m['items']))

    md = os.path.join(OUT, meta['md_filename'])
    n, d = build_md(m, md)
    print('HP用md:', os.path.basename(md), '（%d品・説明文%d件）' % (n, d))


if __name__ == '__main__':
    main()
